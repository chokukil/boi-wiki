from __future__ import annotations

import json
from pathlib import Path

from PIL import Image
from pptx import Presentation


REPO_ROOT = Path(__file__).resolve().parents[1]
DECK_ROOT = REPO_ROOT / "artifacts/science-verifier/deck"


def test_science_evidence_deck_has_exactly_three_grounded_slides() -> None:
    pptx_path = DECK_ROOT / "science-verifier-evidence.pptx"
    manifest = json.loads((DECK_ROOT / "build-manifest.json").read_text())
    presentation = Presentation(pptx_path)

    assert len(presentation.slides) == 3
    assert manifest["slide_count"] == 3
    assert manifest["release_status"] == "release_candidate"
    assert manifest["activation_eligible"] is False
    assert "G5" in " ".join(
        shape.text
        for slide in presentation.slides
        for shape in slide.shapes
        if hasattr(shape, "text")
    )

    for number in (1, 2, 3):
        path = DECK_ROOT / f"rendered/slide-{number}.png"
        with Image.open(path) as image:
            assert image.size == (1600, 900)

    assert (DECK_ROOT / "rendered/whole-deck.png").stat().st_size > 50_000
    assert (DECK_ROOT / "rendered/representative-evidence-slide.png").stat().st_size > 50_000
